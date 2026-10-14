"""Generate original raw-media fixtures for reproducible multimodal validation."""

import argparse
import json
import subprocess
import textwrap
from pathlib import Path

SPEECH = (
    "Welcome to this short lesson on linked lists. The first diagram shows three nodes. "
    "Head points to node seven. Seven points to eleven, and eleven points to nineteen. "
    "The next reference from nineteen is null. "
    "To insert twenty seven after the already known node eleven, first point twenty seven to nineteen. "
    "Then point eleven to twenty seven. These are two constant time reference updates. "
    "If you must first search for eleven, finding that position may inspect every node. "
    "Searching is linear; changing the references is constant time."
)


def create(directory, audio=None):
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import simpleSplit
    from pptx import Presentation
    from pptx.util import Inches, Pt
    import imageio_ffmpeg

    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
    except OSError:
        try:
            font = ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 30
            )
        except OSError:
            font = ImageFont.load_default(size=30)
    scenes = [
        (
            "Follow the next reference",
            [7, 11, 19],
            "The arrows, not the left-to-right placement alone, define the chain.",
        ),
        (
            "Insert after a known node",
            [7, 11, 27, 19],
            "First X.next = A.next. Then A.next = X. Preserve the original successor.",
        ),
        (
            "Finding is not rewiring",
            [7, 11, 27, 19],
            "Searching: O(n) worst case. Rewiring at a known position: O(1).",
        ),
    ]
    images = []
    for index, (title, nodes, caption) in enumerate(scenes):
        image = Image.new("RGB", (1280, 720), "#edf3e9")
        draw = ImageDraw.Draw(image)
        draw.text(
            (65, 50), "LEARNPROOF / ORIGINAL DEMO LESSON", font=font, fill="#607a4f"
        )
        draw.text((65, 135), title, font=font, fill="#183d35")
        for i, value in enumerate(nodes):
            x = 85 + i * 270
            draw.rounded_rectangle(
                (x, 270, x + 180, 390), 15, fill="white", outline="#356455", width=4
            )
            draw.text((x + 62, 305), str(value), font=font, fill="#183d35")
            if i < len(nodes) - 1:
                draw.line((x + 180, 330, x + 265, 330), fill="#356455", width=4)
                draw.polygon(
                    [(x + 265, 330), (x + 248, 320), (x + 248, 340)], fill="#356455"
                )
        draw.text((85, 220), "head", font=font, fill="#536e43")
        draw.text(
            (90 + (len(nodes) - 1) * 270, 420), "next = null", font=font, fill="#536e43"
        )
        draw.multiline_text(
            (65, 545),
            "\n".join(textwrap.wrap(caption, 65)),
            font=font,
            fill="#365540",
            spacing=12,
        )
        path = directory / f"scene-{index}.png"
        image.save(path)
        images.append(path)
    book = directory / "linked-list-evidence.pdf"
    pdf = canvas.Canvas(str(book), pagesize=(612, 792))
    deck = Presentation()
    deck.slide_width = Inches(12.8)
    deck.slide_height = Inches(7.2)
    text = [
        "Pointers identify nodes. A singly linked list follows next references from head to a null reference. Diagram 1 contains values that must be read from the image; they are intentionally absent from this extracted text.",
        "Inserting after an already known node requires two reference updates: X.next = A.next, then A.next = X. Searching for a target value can require n node inspections. Searching is O(n); rewiring is O(1).",
        "The diagrams show a before and after state. Inspect the direction of each arrow to identify the successor and verify that the original tail remains reachable. A changed picture is evidence; an inferred explanation still requires checking.",
    ]
    for i, ((title, _, _), image, paragraph) in enumerate(zip(scenes, images, text), 1):
        pdf.setFont("Helvetica-Bold", 22)
        pdf.drawString(40, 735, title)
        pdf.setFont("Helvetica", 12)
        y = 690
        for line in simpleSplit(paragraph, "Helvetica", 12, 530):
            pdf.drawString(40, y, line)
            y -= 20
        pdf.drawImage(str(image), 40, 250, width=530, height=298)
        pdf.setFont("Helvetica", 10)
        pdf.drawString(40, 65, f"Original LearnProof course fixture / MIT / Page {i}")
        pdf.showPage()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(
            str(image), 0, 0, width=deck.slide_width, height=deck.slide_height
        )
        # The slide is deliberately an embedded figure, not a pre-extracted transcript.
    pdf.save()
    slides = directory / "linked-list-evidence.pptx"
    deck.save(str(slides))
    concat = directory / "scenes.txt"
    concat.write_text(
        "".join(f"file 'scene-{i}.png'\nduration 15\n" for i in range(3))
        + "file 'scene-2.png'\n",
        encoding="utf-8",
    )
    video = directory / "linked-list-lecture.mp4"
    command = [
        imageio_ffmpeg.get_ffmpeg_exe(),
        "-y",
        "-f",
        "concat",
        "-safe",
        "0",
        "-i",
        str(concat),
    ]
    if audio:
        command += ["-i", str(Path(audio).resolve()), "-af", "apad", "-shortest"]
    command += ["-t", "45", "-vf", "fps=12,format=yuv420p", "-c:v", "libx264"]
    if audio:
        command += ["-c:a", "aac"]
    command += [str(video)]
    subprocess.run(command, cwd=directory, check=True, capture_output=True, timeout=120)
    manifest = {
        "license": "MIT",
        "origin": "Original LearnProof authored fixtures",
        "files": [book.name, slides.name, video.name],
        "video_duration_seconds": 45,
        "audio": bool(audio),
        "visual_only_values": [7, 11, 19, 27],
        "note": "Raw PDF, image-based PPTX and MP4; no transcript is passed to ingestion.",
    }
    (directory / "manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )
    (directory / "narration.txt").write_text(SPEECH, encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="tmp/demo-material")
    parser.add_argument("--audio")
    args = parser.parse_args()
    print(json.dumps(create(args.output, args.audio), indent=2))


if __name__ == "__main__":
    main()
