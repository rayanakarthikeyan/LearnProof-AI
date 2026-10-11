"""Original, MIT-licensed course packet and manually reviewed practice items."""
from pathlib import Path

from .ingestion import ingest, register
from .store import DATA, execute, json_value, rows

LESSONS = [
    ('pointers', 'Pointers & references',
     'A pointer stores an address or reference to a node, not a copy of the entire node. '
     'In a singly linked list, a node contains a value and a next reference. '
     'The next reference identifies the following node. A null next reference marks the end. '
     'Two references can point to the same node. Changing one reference does not copy the node.'),
    ('traversal', 'Following the links',
     'Traversal starts at head and follows each next reference in order. To search for a value, '
     'inspect each visited node until the value is found or the list ends. The worst-case search '
     'visits n nodes, so it takes O(n) time. Advancing from the first node to the kth node follows '
     'k-1 next links. Searching for a missing value in a list with n nodes inspects all n nodes.'),
    ('insertion', 'Insert without losing the chain',
     'To insert a new node X after a known node A, first set X.next to A.next, then set A.next '
     'to X. Reversing these updates can lose the original successor. This insertion rewires '
     'two next references. Inserting at the head sets X.next to head, then moves head to X. '
     'No traversal is needed when the insertion position is already known.'),
    ('complexity', 'Find the position. Then rewire.',
     'Insertion after a known node takes O(1) time for the pointer updates. Insertion after '
     'searching for a target value takes O(n) worst-case time: finding the position is linear '
     'and rewiring is constant. Not every linked-list insertion is O(1). For a simple operation '
     'count, total work equals inspected nodes plus pointer updates. Inspecting n nodes and '
     'making two updates gives n+2 counted operations. These counts are a teaching model, '
     'not real processor timings.'),
    ('stacks', 'Stacks: last in, first out',
     'A stack follows LIFO: last in, first out. Push adds to the top and pop removes from the '
     'top. Pushing A, then B, then C and popping once removes C. The empty stack has zero items. '
     'Starting with s items, p successful pushes and q successful pops leaves s+p-q items. '
     'A pop on an empty stack is underflow; it is not a successful pop.'),
    ('queues', 'Queues: first in, first out',
     'A queue follows FIFO: first in, first out. Enqueue adds at the rear and dequeue removes '
     'from the front. Enqueuing A, then B, then C and dequeuing once removes A. '
     'Starting with s items, e successful enqueues and d successful dequeues leaves s+e-d '
     'items. A dequeue from an empty queue is underflow.'),
]


def create_packet(directory: Path):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import simpleSplit
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image, ImageDraw, ImageFont
    directory.mkdir(parents=True, exist_ok=True)
    diagrams = []
    for concept, title, text in LESSONS:
        path = directory / f'{concept}.png'
        im = Image.new('RGB', (1100, 300), '#edf3e9')
        draw = ImageDraw.Draw(im)
        try:
            font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 28)
        except OSError:
            try:
                font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 28)
            except OSError:
                font = ImageFont.load_default(size=28)
        draw.text((35, 25), title, fill='#183d35', font=font)
        for i, label in enumerate(['head / A', 'B', 'C', 'null']):
            x = 35 + i * 270
            draw.rounded_rectangle((x, 125, x+210, 220), 12, fill='#ffffff', outline='#356455', width=3)
            draw.text((x+22, 155), label, fill='#183d35', font=font)
            if i < 3:
                draw.line((x+215, 170, x+260, 170), fill='#356455', width=3)
                draw.polygon([(x+260,170),(x+250,162),(x+250,178)], fill='#356455')
        im.save(path)
        diagrams.append(path)
    pdf = directory / 'LearnProof-course-notes.pdf'
    doc = canvas.Canvas(str(pdf), pagesize=(612, 792))
    deck = Presentation()
    for index, ((concept, title, text), diagram) in enumerate(zip(LESSONS, diagrams), 1):
        doc.setFillColorRGB(.08, .22, .18)
        doc.setFont('Helvetica-Bold', 11)
        doc.drawString(48, 738, 'LEARNPROOF / DATA STRUCTURES FOUNDATIONS')
        doc.setFont('Helvetica-Bold', 23)
        doc.drawString(48, 685, title)
        doc.setFont('Helvetica', 12)
        y = 638
        for line in simpleSplit(text, 'Helvetica', 12, 500):
            doc.drawString(48, y, line)
            y -= 21
        doc.drawImage(str(diagram), 48, 185, width=516, height=141)
        doc.setFont('Helvetica', 10)
        doc.drawString(48, 100, 'Authored demo material. Diagram shows the same linked-list structure on each page.')
        doc.drawString(48, 68, f'Original LearnProof sample / MIT licensed / Page {index}')
        doc.showPage()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(.5), Inches(.4), Inches(9), Inches(.8))
        title_box.text = title
        title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        body = slide.shapes.add_textbox(Inches(.5), Inches(1.4), Inches(9), Inches(2.4))
        body.text_frame.word_wrap = True
        body.text = text
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        slide.shapes.add_picture(str(diagram), Inches(.5), Inches(4.4), width=Inches(9))
    doc.save()
    pptx = directory / 'LearnProof-lecture-slides.pptx'
    deck.save(str(pptx))
    return pdf, pptx


def seed():
    if rows('SELECT id FROM sources LIMIT 1'):
        return
    pdf, pptx = create_packet(DATA / 'sample')
    pdf_id = register(pdf, name='Foundations · authored course notes')
    ingest(pdf_id, use_ai=False)
    slide_id = register(pptx, name='Foundations · authored lecture slides')
    ingest(slide_id, use_ai=False)
    unit_rows = rows('SELECT id,page FROM units WHERE source_id=?', (pdf_id,))
    unit_by_concept = {LESSONS[r['page']-1][0]: r['id'] for r in unit_rows}
    if len(unit_by_concept) != len(LESSONS):
        return
    bank = [
      ('pointers','mcq','What does a pointer store?', ['A reference/address', 'A copy of the entire list', 'The list length', 'An operation count'], 'A reference/address', None, 'A pointer stores a reference or address, not a copy.'),
      ('pointers','short','What does a null next reference indicate?', [], 'The end of the list', ['end','end of list','the end of the list','no following node'], 'A null next reference marks the end.'),
      ('pointers','mcq','Can two references point to the same node?', ['Yes', 'No'], 'Yes', None, 'Two references can identify the same node.'),
      ('traversal','mcq','What is the worst-case cost of searching a singly linked list?', ['O(1)', 'O(n)', 'O(log n)', 'O(n squared)'], 'O(n)', None, 'Search may inspect every node, so worst-case time is linear.'),
      ('traversal','numerical','How many next links must be followed to move from the first to the sixth node?', [], '5', None, 'Reaching the kth node follows k-1 links; 6-1 = 5.'),
      ('traversal','numerical','A value is absent from a 12-node list. How many nodes must a worst-case search inspect?', [], '12', None, 'An absent-value search inspects all n nodes.'),
      ('traversal','short','Which reference do you repeatedly follow to traverse a singly linked list?', [], 'next', ['next','next pointer','the next reference','next reference'], 'Traversal follows each next reference.'),
      ('insertion','mcq','When inserting X after A, which update must happen first?', ['X.next = A.next', 'A.next = X', 'head = null', 'Delete A'], 'X.next = A.next', None, 'Save the old successor into X.next before changing A.next.'),
      ('insertion','numerical','How many next-reference updates insert X after an already known node A?', [], '2', None, 'Set X.next to A.next, then A.next to X: two updates.'),
      ('insertion','short','After setting X.next = A.next, what reference update completes insertion after A?', [], 'A.next = X', ['a.next=x','a.next = x','a->next=x','a->next = x'], 'The second update connects A to X.'),
      ('complexity','mcq','Is every linked-list insertion O(1), including finding a target value first?', ['Yes, all insertion work is O(1)', 'No, finding the target can take O(n)'], 'No, finding the target can take O(n)', None, 'Rewiring is constant; searching for the position can be linear.'),
      ('complexity','mcq','A reference to the insertion node is already provided. What is the rewiring cost?', ['O(1)', 'O(n)', 'O(log n)'], 'O(1)', None, 'No search is needed when the insertion node is known.'),
      ('complexity','numerical','Using the teaching operation-count model, 10 inspected nodes plus two pointer updates totals how many operations?', [], '12', None, 'The model counts inspected nodes plus updates: 10+2 = 12.'),
      ('complexity','short','What is the worst-case time to search for an insertion position and then rewire?', [], 'O(n)', ['o(n)','linear','linear time','o(n) time'], 'Linear search plus constant rewiring is O(n).'),
      ('stacks','mcq','Push A, B, C into an empty stack. Which item is popped first?', ['A','B','C'], 'C', None, 'Last in, first out removes C.'),
      ('stacks','short','What ordering rule does a stack follow?', [], 'LIFO', ['lifo','last in first out','last-in, first-out','last in, first out'], 'A stack is last in, first out.'),
      ('stacks','numerical','A stack starts with 3 items, then has 4 successful pushes and 2 successful pops. How many items remain?', [], '5', None, 'Use s+p-q: 3+4-2 = 5.'),
      ('queues','mcq','Enqueue A, B, C into an empty queue. Which item is dequeued first?', ['A','B','C'], 'A', None, 'First in, first out removes A.'),
      ('queues','short','What ordering rule does a queue follow?', [], 'FIFO', ['fifo','first in first out','first-in, first-out','first in, first out'], 'A queue is first in, first out.'),
      ('queues','numerical','A queue starts with 2 items, then has 5 successful enqueues and 3 successful dequeues. How many remain?', [], '4', None, 'Use s+e-d: 2+5-3 = 4.'),
    ]
    for index, (concept,kind,prompt,options,answer,aliases,feedback) in enumerate(bank):
        execute('INSERT INTO questions VALUES (?,?,?,?,?,?,?,?,?,?)',
                (f'authored-{index+1}', 'data-structures', concept, kind, 1 if concept != 'complexity' else 2,
                 prompt, json_value({'options':options,'answer':answer,'aliases':aliases or [],
                                     'feedback':feedback,'tolerance':0,'origin':'authored-demo'}),
                 json_value([unit_by_concept[concept]]), 'author-reviewed', prompt.lower()))
